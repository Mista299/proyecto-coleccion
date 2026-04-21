import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Cargar archivos TSV ────────────────────────────────────────────────────────
occurrence = pd.read_csv("occurrence.txt", sep="\t", encoding="utf-8", low_memory=False)
emof       = pd.read_csv("extendedmeasurementorfact.txt", sep="\t", encoding="utf-8", low_memory=False)

occurrence.columns = occurrence.columns.str.strip()
emof.columns       = emof.columns.str.strip()

# ── Pivotar extensión de mediciones ───────────────────────────────────────────
emof_pivot = (
    emof[emof["measurementType"].notna()]
    .pivot_table(index="occurrenceID", columns="measurementType",
                 values="measurementValue", aggfunc="first")
    .reset_index()
)
emof_pivot.columns.name = None
df = occurrence.merge(emof_pivot, on="occurrenceID", how="left")

# ── Tipos de datos ─────────────────────────────────────────────────────────────
int_cols   = ["year", "month", "day", "individualCount",
              "minimumElevationInMeters", "maximumElevationInMeters",
              "coordinateUncertaintyInMeters"]
float_cols = ["decimalLatitude", "decimalLongitude",
              "organismQuantity", "coordinatePrecision"]

for c in int_cols + float_cols:
    if c in df.columns:
        df[c] = pd.to_numeric(df[c], errors="coerce")

for c in ["eventDate", "dateIdentified"]:
    if c in df.columns:
        df[c] = pd.to_datetime(df[c], errors="coerce")

# ── Helper: columna si existe, si no pd.NA ─────────────────────────────────────
def col(name):
    return df[name] if name in df.columns else pd.NA

# ══════════════════════════════════════════════════════════════════════════════
# TABLAS DARWIN CORE  ─ https://dwc.tdwg.org/terms/
# ══════════════════════════════════════════════════════════════════════════════

# 1. REGISTRO (Occurrence)
registro = pd.DataFrame({
    "occurrenceID":          col("occurrenceID"),
    "basisOfRecord":         col("basisOfRecord"),
    "type":                  col("type"),
    "institutionCode":       col("institutionCode"),
    "collectionCode":        col("collectionCode"),
    "institutionID":         col("institutionID"),
    "collectionID":          col("collectionID"),
    "catalogNumber":         col("catalogNumber"),
    "occurrenceStatus":      col("occurrenceStatus"),
    "preparations":          col("preparations"),
    "disposition":           col("disposition"),
    "individualCount":       col("individualCount"),
    "organismQuantity":      col("organismQuantity"),
    "organismQuantityType":  col("organismQuantityType"),
    "sex":                   col("sex"),
    "lifeStage":             col("lifeStage"),
    "reproductiveCondition": col("reproductiveCondition"),
    "behavior":              col("behavior"),
    "occurrenceRemarks":     col("occurrenceRemarks"),
})
registro.columns = [
    "ID Registro", "Base del registro", "Tipo",
    "Código institución", "Código colección",
    "ID institución", "ID colección", "N° catálogo",
    "Estado del registro", "Preparaciones", "Disposición",
    "Cantidad individuos", "Cantidad organismo", "Tipo cantidad organismo",
    "Sexo", "Etapa vital", "Condición reproductiva",
    "Comportamiento", "Observaciones",
]

# 2. EVENTO (Event)
evento = pd.DataFrame({
    "occurrenceID":      col("occurrenceID"),
    "eventID":           col("eventID"),
    "eventDate":         col("eventDate"),
    "year":              col("year"),
    "month":             col("month"),
    "day":               col("day"),
    "startDayOfYear":    col("startDayOfYear"),
    "endDayOfYear":      col("endDayOfYear"),
    "verbatimEventDate": col("verbatimEventDate"),
    "samplingProtocol":  col("samplingProtocol"),
    "sampleSizeValue":   col("sampleSizeValue"),
    "sampleSizeUnit":    col("sampleSizeUnit"),
    "samplingEffort":    col("samplingEffort"),
    "fieldNotes":        col("fieldNotes"),
    "recordedBy":        col("recordedBy"),
    "recordNumber":      col("recordNumber"),
    "eventRemarks":      col("eventRemarks"),
})
evento.columns = [
    "ID Registro", "ID Evento", "Fecha", "Año", "Mes", "Día",
    "Día inicio año", "Día fin año", "Fecha verbatim",
    "Protocolo muestreo", "Tamaño muestra", "Unidad muestra",
    "Esfuerzo muestreo", "Notas de campo",
    "Colectado por", "N° colecta", "Observaciones evento",
]

# 3. UBICACIÓN (Location)
ubicacion = pd.DataFrame({
    "occurrenceID":                  col("occurrenceID"),
    "continent":                     col("continent"),
    "country":                       col("country"),
    "countryCode":                   col("countryCode"),
    "stateProvince":                 col("stateProvince"),
    "county":                        col("county"),
    "municipality":                  col("municipality"),
    "locality":                      col("locality"),
    "waterBody":                     col("waterBody"),
    "islandGroup":                   col("islandGroup"),
    "island":                        col("island"),
    "verbatimLocality":              col("verbatimLocality"),
    "minimumElevationInMeters":      col("minimumElevationInMeters"),
    "maximumElevationInMeters":      col("maximumElevationInMeters"),
    "minimumDepthInMeters":          col("minimumDepthInMeters"),
    "maximumDepthInMeters":          col("maximumDepthInMeters"),
    "decimalLatitude":               col("decimalLatitude"),
    "decimalLongitude":              col("decimalLongitude"),
    "geodeticDatum":                 col("geodeticDatum"),
    "coordinateUncertaintyInMeters": col("coordinateUncertaintyInMeters"),
    "coordinatePrecision":           col("coordinatePrecision"),
    "verbatimCoordinates":           col("verbatimCoordinates"),
    "georeferenceRemarks":           col("georeferenceRemarks"),
    "habitat":                       col("habitat"),
    "locationRemarks":               col("locationRemarks"),
})
ubicacion.columns = [
    "ID Registro", "Continente", "País", "Código país",
    "Departamento", "Municipio", "Corregimiento", "Localidad",
    "Cuerpo de agua", "Grupo de islas", "Isla", "Localidad verbatim",
    "Altitud mín (m)", "Altitud máx (m)",
    "Profundidad mín (m)", "Profundidad máx (m)",
    "Latitud", "Longitud", "Datum",
    "Incertidumbre coord (m)", "Precisión coord",
    "Coordenadas verbatim", "Observaciones georeferencia",
    "Hábitat", "Observaciones localidad",
]

# 4. TAXONOMÍA (Taxon)
taxonomia = pd.DataFrame({
    "occurrenceID":             col("occurrenceID"),
    "scientificNameID":         col("scientificNameID"),
    "scientificName":           col("scientificName"),
    "acceptedNameUsage":        col("acceptedNameUsage"),
    "acceptedNameUsageID":      col("acceptedNameUsageID"),
    "originalNameUsage":        col("originalNameUsage"),
    "scientificNameAuthorship": col("scientificNameAuthorship"),
    "taxonRank":                col("taxonRank"),
    "taxonomicStatus":          col("taxonomicStatus"),
    "kingdom":                  col("kingdom"),
    "phylum":                   col("phylum"),
    "class":                    col("class"),
    "order":                    col("order"),
    "family":                   col("family"),
    "subfamily":                col("subfamily"),
    "genus":                    col("genus"),
    "subgenus":                 col("subgenus"),
    "specificEpithet":          col("specificEpithet"),
    "infraspecificEpithet":     col("infraspecificEpithet"),
    "vernacularName":           col("vernacularName"),
    "nomenclaturalCode":        col("nomenclaturalCode"),
    "taxonRemarks":             col("taxonRemarks"),
})
taxonomia.columns = [
    "ID Registro", "ID Nombre científico",
    "Nombre científico", "Nombre aceptado", "ID Nombre aceptado",
    "Nombre original", "Autoría", "Rango taxonómico", "Estado taxonómico",
    "Reino", "Filo", "Clase", "Orden", "Familia", "Subfamilia",
    "Género", "Subgénero", "Epíteto específico", "Epíteto infraespecífico",
    "Nombre común", "Código nomenclatural", "Observaciones taxonómicas",
]

# 5. IDENTIFICACIÓN (Identification)
identificacion = pd.DataFrame({
    "occurrenceID":                     col("occurrenceID"),
    "identifiedBy":                     col("identifiedBy"),
    "dateIdentified":                   col("dateIdentified"),
    "identificationQualifier":          col("identificationQualifier"),
    "typeStatus":                       col("typeStatus"),
    "identificationRemarks":            col("identificationRemarks"),
    "previousIdentifications":          col("previousIdentifications"),
    "identificationVerificationStatus": col("identificationVerificationStatus"),
    "identificationReferences":         col("identificationReferences"),
})
identificacion.columns = [
    "ID Registro", "Identificado por", "Fecha identificación",
    "Calificador identificación", "Tipo nomenclatural",
    "Observaciones identificación", "Identificaciones previas",
    "Estado verificación", "Referencias identificación",
]

# ── Diccionario final ──────────────────────────────────────────────────────────
tablas = {
    "Registro":        registro,
    "Evento":          evento,
    "Ubicación":       ubicacion,
    "Taxonomía":       taxonomia,
    "Identificación":  identificacion,
}

# ── Vista rápida en consola ────────────────────────────────────────────────────
ANCHO = 80
for nombre, tabla in tablas.items():
    print(f"\n{'━' * ANCHO}")
    print(f"  {nombre.upper()}  │  {len(tabla):,} registros  │  {len(tabla.columns)} columnas")
    print(f"{'━' * ANCHO}")
    preview = tabla.head(3).copy()
    for c in preview.select_dtypes(include="object").columns:
        preview[c] = preview[c].astype(str).str[:30]
    print(preview.to_string(index=False, max_colwidth=30))
print(f"\n{'━' * ANCHO}\n")

# ── Exportar a Excel ───────────────────────────────────────────────────────────
RUTA = "coleccion_ciua.xlsx"

with pd.ExcelWriter(RUTA, engine="openpyxl", datetime_format="YYYY-MM-DD") as writer:
    for nombre, tabla in tablas.items():
        tabla.to_excel(writer, sheet_name=nombre, index=False)

# ── Formato profesional ────────────────────────────────────────────────────────
wb = load_workbook(RUTA)

COLORES = {
    "Registro":       {"header": "1F4E79", "fila_par": "D6E4F0"},
    "Evento":         {"header": "4A235A", "fila_par": "E8DAEF"},
    "Ubicación":      {"header": "145A32", "fila_par": "D5F5E3"},
    "Taxonomía":      {"header": "1A5276", "fila_par": "D4E6F1"},
    "Identificación": {"header": "6E2F0A", "fila_par": "FAE5D3"},
}

borde_fino = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

for nombre in tablas:
    ws = wb[nombre]
    color_header = COLORES[nombre]["header"]
    color_par    = COLORES[nombre]["fila_par"]
    max_col = ws.max_column
    max_row = ws.max_row

    for col_idx in range(1, max_col + 1):
        c = ws.cell(row=1, column=col_idx)
        c.font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill      = PatternFill("solid", fgColor=color_header)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = borde_fino
    ws.row_dimensions[1].height = 32

    for row_idx in range(2, max_row + 1):
        for col_idx in range(1, max_col + 1):
            c = ws.cell(row=row_idx, column=col_idx)
            if row_idx % 2 == 0:
                c.fill = PatternFill("solid", fgColor=color_par)
            c.font      = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="center", wrap_text=False)
            c.border    = borde_fino
        ws.row_dimensions[row_idx].height = 16

    for col_idx in range(1, max_col + 1):
        letra   = get_column_letter(col_idx)
        max_len = len(str(ws.cell(row=1, column=col_idx).value or ""))
        for row_idx in range(2, min(max_row + 1, 50)):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[letra].width = min(max_len + 4, 45)

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.tabColor = color_header

wb.save(RUTA)
print(f"✅  Archivo exportado: {RUTA}")